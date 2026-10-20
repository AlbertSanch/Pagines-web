"""Descarga la base de datos del Sistema Estatal de Índices de Alquiler de Vivienda (SERPAVI)
del Ministerio de Vivienda y Agenda Urbana: renta de alquiler declarada a Hacienda (mediana y
percentiles en €/mes y €/m²) por municipio, distrito y sección censal, años 2011-2024.

El Excel (unos 70 MB) se guarda en rentamedia/datos-alquiler/ (no se sube a GitHub). Con
--explorar, además escribe rentamedia/data/alquiler_explorar.txt con las hojas, las cabeceras y
unas filas de ejemplo, para ver la estructura antes de procesarlo.

Uso: python rentamedia/scripts/descargar_alquiler.py [--explorar]
"""
import re
import sys
import time
import urllib.parse
import urllib.request
from pathlib import Path

PAGINA = "https://www.mivau.gob.es/vivienda/alquila-bien-es-tu-derecho/serpavi"
RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "datos-alquiler"
CABECERAS = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36",
             "Accept-Language": "es-ES,es;q=0.9"}


def bajar(url, timeout=300):
    for intento in range(4):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers=CABECERAS), timeout=timeout) as r:
                return r.read()
        except Exception as e:
            print(f"{url}: {e} (intento {intento + 1})", file=sys.stderr)
            time.sleep(15 * (intento + 1))
    sys.exit(f"ERROR: no se ha podido descargar {url}")


def main():
    SALIDA.mkdir(exist_ok=True)
    html = bajar(PAGINA).decode("utf-8", "replace")
    enlaces = [urllib.parse.urljoin(PAGINA, h) for h in re.findall(r'href="([^"]+\.xlsx?)"', html, flags=re.I)]
    print("Excel enlazados en la página:\n  " + "\n  ".join(enlaces))
    bd = [u for u in enlaces if re.search(r"bd|base", urllib.parse.unquote(u), re.I)] or enlaces
    if not bd:
        sys.exit("ERROR: la página del SERPAVI no enlaza ningún Excel (¿ha cambiado la web del Ministerio?)")
    url = bd[0]
    datos = bajar(url)
    if datos[:2] != b"PK":
        sys.exit(f"ERROR: {url} no es un Excel (empieza por {datos[:60]!r})")
    destino = SALIDA / "serpavi.xlsx"
    destino.write_bytes(datos)
    (SALIDA / "origen.txt").write_text(url + "\n", encoding="utf-8")
    print(f"OK {destino} ({len(datos) // 1024 // 1024} MB) desde {url}")

    if "--explorar" in sys.argv:
        import openpyxl
        libro = openpyxl.load_workbook(destino, read_only=True, data_only=True)
        lineas = [f"Fuente: {url}", f"Hojas: {libro.sheetnames}", ""]
        for hoja in libro.worksheets:
            lineas.append(f"=== {hoja.title} (filas: {hoja.max_row}, columnas: {hoja.max_column})")
            for i, fila in enumerate(hoja.iter_rows(values_only=True)):
                if i >= 12:
                    break
                lineas.append(" | ".join("" if v is None else str(v) for v in fila))
            lineas.append("")
        (RAIZ / "data" / "alquiler_explorar.txt").write_text("\n".join(lineas), encoding="utf-8")
        print("\n".join(lineas[:60]))


if __name__ == "__main__":
    main()
