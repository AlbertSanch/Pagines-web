"""Procesa la base de datos del SERPAVI (rentamedia/datos-alquiler/serpavi.xlsx, de descargar_alquiler.py)
y escribe rentamedia/data/alquiler.json, con el alquiler declarado a Hacienda por territorio:

  campos: n (viviendas alquiladas), eur (mediana €/mes), eur25 y eur75 (percentiles 25 y 75 €/mes),
          m2 (mediana €/m² al mes), sup (mediana de superficie en m²)
  - ccaa, prov y mun: serie de todos los años de pisos («vc»: vivienda colectiva) y último año de
    casas («vu»: vivienda unifamiliar).
  - dis y sec: último año de pisos y la mediana €/mes de cinco años antes (para ver cuánto ha subido).

También guarda la hoja de metadatos del Ministerio en data/alquiler_metadatos.txt.

Uso: python rentamedia/scripts/procesar_alquiler.py
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
ORIGEN = RAIZ / "datos-alquiler" / "serpavi.xlsx"
SALIDA = RAIZ / "data" / "alquiler.json"
CAMPOS = ["n", "eur", "eur25", "eur75", "m2", "sup"]
# columna del Ministerio -> (campo, tipología)
COLUMNA = re.compile(r"^(BI_ALVHEPCO_T|ALQTBID12_(M|25|75)_|ALQM2_LV_M_|SLVM2_M_)(VC|VU)_(\d{2})$")
HOJAS = {"CCAA": ("ccaa", "CCAA", "LITCCAA"), "Provincias": ("prov", "CPRO", "LITPRO"), "Municipios": ("mun", "CUMUN", "NMUN"),
         "Distritos": ("dis", "CUDIS", None), "Secciones censales": ("sec", "CUSEC", None)}
MIN_MUNICIPIOS = 6000


def campo(col):
    m = COLUMNA.match(col)
    if not m:
        return None
    pref, pct, tipo, anio = m.groups()
    if pref.startswith("BI_"):
        c = "n"
    elif pref.startswith("ALQTBID12"):
        c = {"M": "eur", "25": "eur25", "75": "eur75"}[pct]
    elif pref.startswith("ALQM2"):
        c = "m2"
    else:
        c = "sup"
    return c, tipo.lower(), 2000 + int(anio)


def redondear(c, v):
    if v in (None, ""):
        return None
    try:
        v = float(v)
    except (TypeError, ValueError):
        return None
    return round(v, 2) if c == "m2" else round(v)


def main():
    if not ORIGEN.exists():
        sys.exit(f"ERROR: falta {ORIGEN}: ejecuta antes descargar_alquiler.py")
    libro = openpyxl.load_workbook(ORIGEN, read_only=True, data_only=True)
    meta = ["\t".join("" if v is None else str(v) for v in fila) for fila in libro["Metadatos"].iter_rows(values_only=True)]
    (RAIZ / "data" / "alquiler_metadatos.txt").write_text("\n".join(meta) + "\n", encoding="utf-8")

    anios = set()
    brutos = {}
    for hoja, (clave, col_id, col_nombre) in HOJAS.items():
        filas = libro[hoja].iter_rows(values_only=True)
        cab = [str(c or "").strip() for c in next(filas)]
        if col_id not in cab:
            sys.exit(f"ERROR: la hoja {hoja} no tiene la columna {col_id} (¿ha cambiado el formato?)")
        i_id = cab.index(col_id)
        i_nombre = cab.index(col_nombre) if col_nombre in cab else None
        mapa = {i: campo(c) for i, c in enumerate(cab) if campo(c)}
        anios |= {a for _, _, a in mapa.values()}
        datos = {}
        for fila in filas:
            cod = fila[i_id]
            if cod in (None, ""):
                continue
            cod = str(cod).strip()
            d = datos.setdefault(cod, {"nombre": str(fila[i_nombre]).strip() if i_nombre is not None and fila[i_nombre] else None})
            for i, (c, tipo, anio) in mapa.items():
                v = redondear(c, fila[i] if i < len(fila) else None)
                if v is not None:
                    d.setdefault(tipo, {}).setdefault(anio, {})[c] = v
        brutos[clave] = datos
        print(f"{hoja}: {len(datos)} territorios")

    anios = sorted(anios)
    ultimo, hace5 = anios[-1], anios[-1] - 5

    def serie(d, tipo):
        por_anio = d.get(tipo, {})
        return [[por_anio.get(a, {}).get(c) for c in CAMPOS] if a in por_anio else None for a in anios]

    def fila(d, tipo, anio):
        x = d.get(tipo, {}).get(anio)
        return [x.get(c) for c in CAMPOS] if x else None

    salida = {"anios": anios, "campos": CAMPOS, "fuente": (RAIZ / "datos-alquiler" / "origen.txt").read_text().strip()
              if (RAIZ / "datos-alquiler" / "origen.txt").exists() else ""}
    for clave in ("ccaa", "prov", "mun"):
        salida[clave] = {cod: {"nombre": d["nombre"], "vc": serie(d, "vc"), "vu": fila(d, "vu", ultimo)}
                         for cod, d in brutos[clave].items() if d.get("vc")}
    for clave in ("dis", "sec"):
        salida[clave] = {}
        for cod, d in brutos[clave].items():
            f = fila(d, "vc", ultimo)
            if f and f[1] is not None:
                antes = d.get("vc", {}).get(hace5, {}).get("eur")
                salida[clave][cod] = f + [antes]
    n_mun = sum(1 for d in salida["mun"].values() if d["vc"][-1] and d["vc"][-1][1])
    if n_mun < MIN_MUNICIPIOS:
        sys.exit(f"ERROR: solo {n_mun} municipios con alquiler en {ultimo} (¿ha cambiado el formato?)")
    SALIDA.write_text(json.dumps(salida, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"alquiler.json: años {anios[0]}-{ultimo}, {n_mun} municipios con dato en {ultimo}, "
          f"{len(salida['dis'])} distritos, {len(salida['sec'])} secciones, {SALIDA.stat().st_size // 1024} KB")


if __name__ == "__main__":
    main()
