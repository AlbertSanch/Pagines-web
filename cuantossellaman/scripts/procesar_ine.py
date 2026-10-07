"""Convierte las hojas de cálculo del INE en JSON intermedios dentro de cuantossellaman/data/.

Entrada (cuantossellaman/datos-ine/, la llena descargar_ine.py):
  - nombres_por_edad_media.xlsx  → todos los nombres con 20 o más personas, por sexo, con edad media
  - nombres_mas_frecuentes.xlsx  → top 100 de España y top 50 por provincia de residencia
  - nombres_por_fecha.xlsx       → top 50 por década de nacimiento en España (y top por provincia)
  - nomnacAA.xlsx                → nombres de los recién nacidos de cada año (top 100 España, top 10 por comunidad)
  - apellidos_frecuencia.xls     → todos los apellidos con 20 o más personas como primer apellido
  - apellidos_mas_frecuentes.xls → top 100 de España y top 50 por provincia de residencia y de nacimiento

Salida (cuantossellaman/data/):
  - nombres.json, provincias.json, decadas.json, bebes.json, apellidos.json

No inventa ni estima nada: los nombres que no están en los ficheros (menos de 20 personas en
España) simplemente no aparecen. Falla con un error claro si el formato del INE cambia.
Uso: python cuantossellaman/scripts/procesar_ine.py
"""
import json
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import xlrd

from tildes import normalizar

RAIZ = Path(__file__).resolve().parent.parent
ENTRADA = RAIZ / "datos-ine"
DATOS = RAIZ / "data"
SEXOS = {"hombres": "H", "mujeres": "M"}


class FormatoINE(Exception):
    pass


def comprobar(cond, msg):
    if not cond:
        raise FormatoINE(msg)


def filas(ruta, hoja):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    comprobar(hoja in wb.sheetnames, f"{ruta.name}: falta la hoja «{hoja}» (hojas: {wb.sheetnames})")
    return [list(r) for r in wb[hoja].iter_rows(values_only=True)]


def fecha_referencia(rows, fichero):
    for r in rows[:5]:
        for c in r:
            m = re.search(r"(\d{2})/(\d{2})/(\d{4})", str(c or ""))
            if m:
                return f"{m[3]}-{m[2]}-{m[1]}"
    raise FormatoINE(f"{fichero}: no encuentro la fecha de referencia (dd/mm/aaaa) en la cabecera")


def entero(x):
    if isinstance(x, (int, float)) and x == x:
        return int(x)
    if isinstance(x, str) and x.strip().replace(".", "").isdigit():
        return int(x.strip().replace(".", ""))
    return None


def limpiar(nombre):
    return " ".join(str(nombre).strip().upper().split())


# --------------------------------------------------------------------------------- Todos los nombres
def leer_todos():
    ruta = ENTRADA / "nombres_por_edad_media.xlsx"
    resultado, referencia = {}, None
    for hoja, sexo in (("Hombres", "H"), ("Mujeres", "M")):
        rows = filas(ruta, hoja)
        referencia = referencia or fecha_referencia(rows, ruta.name)
        cab = next((i for i, r in enumerate(rows[:20]) if r[:3] == ["Orden", "Nombre", "Frecuencia"]), None)
        comprobar(cab is not None, f"{ruta.name}/{hoja}: no encuentro la cabecera «Orden, Nombre, Frecuencia»")
        comprobar("Edad" in str(rows[cab][3]), f"{ruta.name}/{hoja}: la 4.ª columna ya no es la edad media")
        lista = []
        for r in rows[cab + 1:]:
            if r[1] is None:
                continue
            f = entero(r[2])
            comprobar(f is not None and f >= 20, f"{ruta.name}/{hoja}: frecuencia no válida en {r[:4]}")
            lista.append({"n": limpiar(r[1]), "f": f, "r": entero(r[0]),
                          "e": round(float(r[3]), 1) if isinstance(r[3], (int, float)) else None})
        comprobar(len(lista) > 10000, f"{ruta.name}/{hoja}: solo {len(lista)} nombres; esperaba más de 10.000")
        resultado[sexo] = lista
    return referencia, resultado


# --------------------------------------------------------------------------- Tablas «ancha» del INE
def bloques_ancho(rows, fila_cab, fila_col, fin=None):
    """Lee una tabla con un bloque NOMBRE/FRECUENCIA/Por 1.000 por cada provincia o década.

    fila_cab: fila con la etiqueta de cada bloque (provincia o década) en la columna NOMBRE.
    fila_col: fila con los títulos NOMBRE / FRECUENCIA / Por 1.000.
    """
    bloques = {}
    for c, titulo in enumerate(rows[fila_col]):
        if str(titulo or "").strip().upper() != "NOMBRE":
            continue
        etiqueta = rows[fila_cab][c] if c < len(rows[fila_cab]) else None
        comprobar(etiqueta, f"no hay etiqueta encima de la columna {c}")
        lista = []
        for r in rows[fila_col + 1:fin]:
            if c + 1 >= len(r) or r[c] is None:
                if lista:
                    break
                continue
            f = entero(r[c + 1])
            if f is None:
                break
            lista.append([limpiar(r[c]), f])
        bloques[str(etiqueta).strip()] = lista
    return bloques


def leer_provincias():
    ruta = ENTRADA / "nombres_mas_frecuentes.xlsx"
    provincias, referencia = {}, None
    for hoja, sexo in (("Hombres_PROVINCIAdeRESIDENCIA", "H"), ("Mujeres_PROVINCIAdeRESIDENCIA", "M")):
        rows = filas(ruta, hoja)
        referencia = referencia or fecha_referencia(rows, ruta.name)
        fila_col = next(i for i, r in enumerate(rows[:10]) if "NOMBRE" in [str(x).strip() for x in r if x])
        bloques = bloques_ancho(rows, fila_col - 1, fila_col)
        comprobar(len(bloques) == 52, f"{ruta.name}/{hoja}: hay {len(bloques)} provincias, esperaba 52")
        for etiqueta, lista in bloques.items():
            m = re.match(r"(\d{2})\s*-\s*(.+)", etiqueta)
            comprobar(m, f"{ruta.name}/{hoja}: etiqueta de provincia inesperada «{etiqueta}»")
            comprobar(len(lista) >= 20, f"{ruta.name}/{hoja}/{etiqueta}: solo {len(lista)} nombres")
            p = provincias.setdefault(m[1], {"ine": m[2].strip()})
            p[sexo] = lista
    nacional = {}
    for hoja, sexo in (("ESPAÑA_100_hombres", "H"), ("ESPAÑA_100_mujeres", "M")):
        rows = filas(ruta, hoja)
        fila_col = next(i for i, r in enumerate(rows[:10]) if r and str(r[1] or "").strip() == "NOMBRE COMPLETO")
        lista = [[limpiar(r[1]), entero(r[2])] for r in rows[fila_col + 1:] if r[1] and entero(r[2])]
        comprobar(len(lista) >= 90, f"{ruta.name}/{hoja}: solo {len(lista)} nombres")
        nacional[sexo] = lista
    return referencia, provincias, nacional


def leer_decadas():
    ruta = ENTRADA / "nombres_por_fecha.xlsx"
    decadas, referencia = {}, None
    for hoja, sexo in (("ESPAÑA_hombres", "H"), ("ESPAÑA_mujeres", "M")):
        rows = filas(ruta, hoja)
        referencia = referencia or fecha_referencia(rows, ruta.name)
        fila_col = next(i for i, r in enumerate(rows[:10]) if "NOMBRE" in [str(x).strip() for x in r if x])
        bloques = bloques_ancho(rows, fila_col - 1, fila_col)
        comprobar(len(bloques) >= 10, f"{ruta.name}/{hoja}: solo {len(bloques)} décadas")
        for etiqueta, lista in bloques.items():
            if "ANTES" in etiqueta.upper():
                clave, desde, hasta = "antes-1930", None, 1929
            else:
                anios = [int(x) for x in re.findall(r"\b(19\d\d|20\d\d)\b", etiqueta)]
                comprobar(len(anios) == 2, f"{ruta.name}/{hoja}: década inesperada «{etiqueta}»")
                desde, hasta = anios
                clave = str(desde)
            d = decadas.setdefault(clave, {"desde": desde, "hasta": hasta})
            d["hasta"] = max(d["hasta"], hasta) if d.get("hasta") else hasta
            d[sexo] = lista
    return referencia, decadas


def leer_bebes():
    anios = {}
    for ruta in sorted(ENTRADA.glob("nomnac*.xlsx")):
        wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        portada = [c for r in wb[wb.sheetnames[0]].iter_rows(values_only=True) for c in r if c]
        m = next((re.search(r"Año (\d{4})", str(c)) for c in portada if re.search(r"Año (\d{4})", str(c))), None)
        comprobar(m, f"{ruta.name}: no encuentro el año en la portada")
        anio = int(m[1])
        datos = {}
        for hoja in wb.sheetnames[1:]:
            rows = [list(r) for r in wb[hoja].iter_rows(values_only=True)]
            fila = next((i for i, r in enumerate(rows[:6]) if "NIÑOS" in [str(x).strip() for x in r if x]), None)
            comprobar(fila is not None, f"{ruta.name}/{hoja}: no encuentro la cabecera NIÑOS / NIÑAS")
            res = {}
            for etiqueta, sexo in (("NIÑOS", "H"), ("NIÑAS", "M")):
                c = [str(x).strip() if x else "" for x in rows[fila]].index(etiqueta)
                lista, total = [], None
                for r in rows[fila + 1:]:
                    if c >= len(r) or r[c - 1] is None:
                        continue
                    if str(r[c - 1]).strip().upper() == "TOTAL":
                        total = entero(r[c])
                        continue
                    f = entero(r[c])
                    if f:
                        lista.append([limpiar(r[c - 1]), f])
                res[sexo] = {"total": total, "top": lista}
            nombre = str(rows[0][0] or hoja).strip()
            clave = "espana" if hoja.upper() == "TOTAL" else normalizar(hoja).replace(" ", "-")
            datos[clave] = {"nombre": "España" if clave == "espana" else nombre, **res}
        comprobar("espana" in datos and len(datos["espana"]["H"]["top"]) >= 50,
                  f"{ruta.name}: falta el total nacional de bebés")
        anios[str(anio)] = datos
    comprobar(anios, "no hay ficheros de recién nacidos (nomnacAA.xlsx)")
    return anios


# ------------------------------------------------------------------------------------- Apellidos
def filas_xls(ruta, hoja):
    wb = xlrd.open_workbook(ruta)
    comprobar(hoja in wb.sheet_names(), f"{ruta.name}: falta la hoja «{hoja}» (hojas: {wb.sheet_names()})")
    s = wb.sheet_by_name(hoja)
    return [[None if c == "" else c for c in s.row_values(i)] for i in range(s.nrows)]


def cifra_o_nada(x):
    """El INE pone «..» cuando no publica la cifra (por privacidad): se guarda como None."""
    return None if x in (None, "..") else entero(x)


def leer_apellidos():
    ruta = ENTRADA / "apellidos_frecuencia.xls"
    lista, referencia = [], None
    for hoja in ("Apellidos >=100", "Apellidos >=20 y <=99"):
        rows = filas_xls(ruta, hoja)
        referencia = referencia or fecha_referencia(rows, ruta.name)
        cab = next((i for i, r in enumerate(rows[:10]) if r[:2] == ["Orden", "Apellido"]), None)
        comprobar(cab is not None, f"{ruta.name}/{hoja}: no encuentro la cabecera «Orden, Apellido»")
        titulos = [str(x or "") for x in rows[cab - 1]]
        comprobar("1º" in titulos[2] and "2º" in titulos[3] and "Ambos" in titulos[4],
                  f"{ruta.name}/{hoja}: las columnas ya no son 1.er apellido, 2.º apellido y ambos ({titulos[:5]})")
        for r in rows[cab + 1:]:
            if r[1] is None:
                continue
            p1 = entero(r[2])
            comprobar(p1 is not None and p1 >= 20, f"{ruta.name}/{hoja}: primer apellido no válido en {r[:5]}")
            lista.append({"n": limpiar(r[1]), "r": entero(r[0]), "p1": p1, "p2": cifra_o_nada(r[3]), "ambos": cifra_o_nada(r[4])})
    comprobar(len(lista) > 50000, f"{ruta.name}: solo {len(lista)} apellidos; esperaba más de 50.000")

    ruta = ENTRADA / "apellidos_mas_frecuentes.xls"
    rows = filas_xls(ruta, "ESPAÑA_100")
    fila = next(i for i, r in enumerate(rows[:10]) if "PRIMER APELLIDO" in [str(x).strip() for x in r if x])
    top100 = [[limpiar(r[1]), entero(r[2])] for r in rows[fila + 1:] if r[1] and entero(r[2])]
    comprobar(len(top100) >= 90, f"{ruta.name}/ESPAÑA_100: solo {len(top100)} apellidos")
    provincias = {}
    for hoja, clave in (("PROVINCIAS_RESIDENCIA", "res"), ("PROVINCIAS_NACIMIENTO", "nac")):
        rows = filas_xls(ruta, hoja)
        fila = next(i for i, r in enumerate(rows[:10]) if "PRIMER APELLIDO" in [str(x).strip() for x in r if x])
        rows[fila] = ["NOMBRE" if str(x or "").strip() == "PRIMER APELLIDO" else x for x in rows[fila]]
        # La hoja de nacimiento trae además «66 - NACIDOS EN EL EXTRANJERO»: solo usamos las 52 provincias
        bloques = {k: v for k, v in bloques_ancho(rows, fila - 1, fila).items() if re.match(r"(0[1-9]|[1-4]\d|5[0-2])\s*-", k)}
        comprobar(len(bloques) == 52, f"{ruta.name}/{hoja}: hay {len(bloques)} provincias, esperaba 52")
        for etiqueta, l in bloques.items():
            m = re.match(r"(\d{2})\s*-\s*(.+)", etiqueta)
            comprobar(m and len(l) >= 20, f"{ruta.name}/{hoja}: provincia inesperada «{etiqueta}» o con pocos apellidos")
            provincias.setdefault(m[1], {})[clave] = l
    return referencia, lista, top100, provincias


def main():
    try:
        ref_todos, todos = leer_todos()
        ref_prov, provincias, nacional = leer_provincias()
        ref_dec, decadas = leer_decadas()
        bebes = leer_bebes()
        ref_ape, apellidos, top_ape, prov_ape = leer_apellidos()
    except FormatoINE as e:
        sys.exit(f"ERROR: el formato de los ficheros del INE ha cambiado: {e}")


    DATOS.mkdir(exist_ok=True)
    meta = {"generado": datetime.now().strftime("%Y-%m-%d"), "fuente": "INE, Estadística de nombres y apellidos (Censos de población anuales)"}
    escribir("nombres.json", {**meta, "referencia": ref_todos, "H": todos["H"], "M": todos["M"], "top100": nacional})
    escribir("provincias.json", {**meta, "referencia": ref_prov, "provincias": provincias})
    escribir("decadas.json", {**meta, "referencia": ref_dec, "decadas": decadas})
    escribir("apellidos.json", {**meta, "referencia": ref_ape, "lista": apellidos, "top100": top_ape, "provincias": prov_ape})
    escribir("bebes.json", {**meta, "fuente": "INE, Estadística de nacimientos (nombres de los recién nacidos)", "anios": bebes})
    print(f"Nombres: {len(todos['H'])} de hombre y {len(todos['M'])} de mujer (datos a {ref_todos}); "
          f"{len(provincias)} provincias; {len(decadas)} décadas; bebés {min(bebes)}–{max(bebes)}; "
          f"{len(apellidos)} apellidos (datos a {ref_ape})")


def escribir(nombre, datos):
    (DATOS / nombre).write_text(json.dumps(datos, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


if __name__ == "__main__":
    main()
